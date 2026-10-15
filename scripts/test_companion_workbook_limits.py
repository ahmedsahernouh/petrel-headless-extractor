# Copyright 2026 Ahmed Saher Nouh
# SPDX-License-Identifier: Apache-2.0
# https://saherlabs.dev/ | https://github.com/ahmedsahernouh/petrel-headless-extractor
"""Workbook limits must retain originals and never publish partial conversions."""
import csv
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import portable_petrel_companion_extract as c


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='FV workbook ')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root/'input.xlsx'
        self.out = self.root/'converted'

    def fixture(self, sheets=None):
        wb = Workbook(); wb.remove(wb.active)
        for title, rows in (sheets or [('Data', [['A', 'B'], [1, None], ['quoted, text', 2.5]])]):
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append(row)
        wb.save(self.source); wb.close()
        return self.source.read_bytes()

    def test_csv_bytes_unchanged_and_original_retained(self):
        original = self.fixture()
        output, summary = c.convert_workbook(self.source, self.out)
        expected = io.StringIO(newline=''); writer = csv.writer(expected)
        writer.writerows([['A', 'B'], [1, ''], ['quoted, text', 2.5]])
        self.assertEqual((self.out/'Data.csv').read_bytes(), expected.getvalue().encode('utf-8'))
        self.assertEqual(summary['sheets'][0]['rows'], 3)
        self.assertEqual(len(output), 2)
        self.assertEqual(self.source.read_bytes(), original)

    def test_cell_limit_across_sheets_removes_all_derivatives_closes_workbook(self):
        original = self.fixture([('First', [[1, 2]]), ('Second', [[3, 4]])])
        wb = load_workbook(self.source, read_only=True, data_only=True)
        with patch('openpyxl.load_workbook', return_value=wb), patch.object(wb, 'close', wraps=wb.close) as closed:
            with self.assertRaises(c.WorkbookLimitError):
                c.convert_workbook(self.source, self.out, max_cells=3)
            closed.assert_called_once()
        self.assertFalse(self.out.exists())
        self.assertEqual(self.source.read_bytes(), original)

    def test_cooperative_time_limit_and_diagnostic(self):
        self.fixture()
        with patch.object(c.time, 'monotonic', side_effect=[0, 0, 301, 301]), \
             patch.object(c, 'workbook_event') as events:
            with self.assertRaises(c.WorkbookLimitError):
                c.convert_workbook(self.source, self.out)
        self.assertFalse(self.out.exists())
        limited = [call for call in events.call_args_list if call.args[0]=='companion_workbook_limited']
        self.assertEqual(len(limited), 1)
        self.assertEqual(limited[0].kwargs['rows'], 1)
        self.assertEqual(limited[0].kwargs['cells'], 2)

    def test_progress_and_telemetry_failure_do_not_change_data(self):
        self.fixture()
        with patch.object(c.time, 'monotonic', side_effect=[0, 6, 7, 8, 9]), \
             patch.object(c, 'workbook_event') as events:
            c.convert_workbook(self.source, self.out, max_seconds=100)
        self.assertTrue(any(call.args[0]=='companion_workbook_progress' for call in events.call_args_list))
        with patch('geoviewer_diagnostics.event', side_effect=OSError('log denied')):
            c.convert_workbook(self.source, self.root/'second')
        self.assertEqual((self.out/'Data.csv').read_bytes(), (self.root/'second/Data.csv').read_bytes())

    def test_existing_directory_is_never_overwritten(self):
        self.fixture(); self.out.mkdir(); existing=self.out/'Data.csv'; existing.write_text('keep')
        with self.assertRaises(FileExistsError):
            c.convert_workbook(self.source, self.out)
        self.assertEqual(existing.read_text(), 'keep')

    def test_writer_failure_closes_workbook_and_removes_derivatives(self):
        original=self.fixture()
        wb=load_workbook(self.source,read_only=True,data_only=True)
        with patch('openpyxl.load_workbook',return_value=wb), \
             patch.object(wb,'close',wraps=wb.close) as closed, \
             patch.object(c.csv,'writer',side_effect=OSError('write failure')):
            with self.assertRaisesRegex(OSError,'write failure'):
                c.convert_workbook(self.source,self.out)
            closed.assert_called_once()
        self.assertFalse(self.out.exists())
        self.assertEqual(self.source.read_bytes(),original)

    def test_sheet_name_collisions_keep_both_and_empty_is_not_exported(self):
        self.fixture([('A B', [[1]]), ('A_B', [[2]]), ('Empty', [])])
        output, summary = c.convert_workbook(self.source, self.out)
        csvs = [p for p in output if p.suffix=='.csv']
        self.assertEqual(len(csvs), 2)
        self.assertEqual({p.read_text().strip() for p in csvs}, {'1','2'})
        self.assertEqual(summary['sheets'][-1]['status'], 'empty_not_exported')

    def test_cli_preserves_and_indexes_workbook_limit_then_continues(self):
        original=self.fixture()
        project=self.root/'Fixture.pet'; project.write_text('fixture')
        (self.root/'Fixture.ptd').mkdir(); (self.root/'notes.txt').write_text('retain notes')
        with tempfile.TemporaryDirectory(prefix='FV output ') as directory:
            package=Path(directory); (package/'00_manifest').mkdir()
            (package/'00_manifest/export_manifest.csv').write_text('export_file,sha256\n')
            argv=['companion','--project-file',str(project),'--export-package',str(package),
                  '--project-name','Fixture','--workbook-max-cells','1']
            with patch.object(sys,'argv',argv), patch('builtins.print'):
                self.assertEqual(c.main(),0)
            with (package/'00_manifest/companion_source_inventory.csv').open(encoding='utf-8',newline='') as stream:
                rows=list(csv.DictReader(stream))
            row=next(r for r in rows if r['source_relative_path']=='input.xlsx')
            self.assertEqual(row['status'],'preserved_conversion_limited')
            self.assertEqual(row['converted_files'],'')
            self.assertEqual((package/row['preserved_file']).read_bytes(),original)
            report=json.loads((package/'07_workflows_reports/portable_extractor/companion_capability_report.json').read_text())
            self.assertEqual(report['counts']['conversion_limited'],1)
            self.assertEqual(report['counts']['converted_source_files'],0)
            self.assertTrue(any(r['source_relative_path']=='notes.txt' for r in rows))
            self.assertIn('preserved_conversion_limited',(package/'99_unexported_or_manual/portable_unsupported_inventory.csv').read_text())
        self.assertEqual(self.source.read_bytes(),original)


if __name__=='__main__':
    unittest.main()
