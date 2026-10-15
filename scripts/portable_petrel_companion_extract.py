#!/usr/bin/env python3
# Copyright 2026 Ahmed Saher Nouh
# SPDX-License-Identifier: Apache-2.0
# Petrel Headless Extractor - Ahmed Saher Nouh / SaherLabs
# Website: https://saherlabs.dev/
# GitHub: https://github.com/ahmedsahernouh
# Repository: https://github.com/ahmedsahernouh/petrel-headless-extractor

"""Inventory, preserve, and convert open companion files beside a Petrel project.

The source project is read-only. Native .pet/.ptd stores are deliberately
excluded because the portable wrapper handles them with the guarded native
exporter. Optional converter failures are recorded per file and do not turn an
unsupported proprietary format into a false successful conversion.
"""

from __future__ import annotations

import argparse
from contextlib import ExitStack
import csv
import datetime as dt
import hashlib
import petrel_progress as progress
import json
from itertools import zip_longest
import os
import re
import shutil
import shlex
import sys
import time
import tempfile
from pathlib import Path
from typing import Any


TEXT_SUFFIXES = {".txt", ".prn", ".csv", ".tsv", ".asc", ".ascii", ".dat", ".prj", ".xml", ".json", ".md"}
PRESERVE_SUFFIXES = TEXT_SUFFIXES | {".las", ".xlsx", ".xlsm", ".xls", ".shp", ".shx", ".dbf", ".sbn", ".sbx", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".sgy", ".segy", ".zgy", ".resqml", ".epc"}
TEXT_SAMPLE_BYTES = 65536
HEADER_SAMPLE_BYTES = 262144
TEXT_PROFILE_BYTES = 1048576
WORKBOOK_BUFFER_BYTES = 1048576
WORKBOOK_MAX_SECONDS = 300.0
WORKBOOK_MAX_CELLS = 5_000_000
INVENTORY_FIELDS = (
    "source_relative_path", "category", "extension", "size_bytes", "sha256",
    "text_line_count", "text_nonempty_line_count", "text_profile_scope",
    "text_sample_bytes", "text_sample_line_count", "likely_delimiter",
    "first_nonempty_line", "status", "preserved_file", "converted_files",
    "conversion_gap", "error",
)


class WorkbookLimitError(RuntimeError):
    """Optional workbook conversion stopped; its source copy remains usable."""


def workbook_event(kind: str, **fields: Any) -> None:
    try:
        from geoviewer_diagnostics import event
        event(kind, **fields)
    except Exception:
        pass  # Diagnostics must not turn a usable companion into a failure.


def sha256(path: Path) -> str:
    return progress.hash_file(path)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def rel(base: Path, path: Path) -> str:
    return os.path.relpath(path, base).replace("/", "\\")


def is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def looks_text(path: Path) -> bool:
    try:
        with path.open("rb") as handle:
            sample = handle.read(TEXT_SAMPLE_BYTES)
    except OSError:
        return False
    if not sample:
        return True
    if b"\x00" in sample:
        return False
    printable = sum(byte in b"\t\n\r" or 32 <= byte <= 126 or byte >= 128 for byte in sample)
    return printable / len(sample) >= 0.90


def classify(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".las": "las_well_data",
        ".xlsx": "excel_workbook",
        ".xlsm": "excel_workbook",
        ".xls": "legacy_excel_workbook",
        ".shp": "gis_shapefile_geometry",
        ".shx": "gis_shapefile_index",
        ".dbf": "gis_shapefile_attributes",
        ".prj": "coordinate_reference_system",
        ".sgy": "segy_seismic",
        ".segy": "segy_seismic",
        ".zgy": "zgy_seismic",
        ".resqml": "resqml_data",
        ".epc": "resqml_epc_package",
        ".pdf": "document_pdf",
        ".xml": "xml_metadata",
        ".csv": "delimited_ascii",
        ".tsv": "delimited_ascii",
        ".txt": "text_ascii",
        ".prn": "text_ascii",
        ".asc": "text_ascii",
        ".ascii": "text_ascii",
        ".dat": "ambiguous_dat_or_zmap",
    }.get(suffix, "text_ascii_no_extension" if not suffix and looks_text(path) else "unsupported_or_unknown")


def profile_text(path: Path) -> dict[str, Any]:
    # Bound both the read and the temporary line lists, including giant files
    # with no line breaks. Counts from a prefix must never look like full totals.
    with path.open("rb") as handle:
        data = handle.read(TEXT_PROFILE_BYTES + 1)
    truncated = len(data) > TEXT_PROFILE_BYTES
    text = data[:TEXT_PROFILE_BYTES].decode("utf-8", errors="replace")
    lines = text.splitlines()
    if truncated and not text.endswith(("\n", "\r")):
        lines = lines[:-1]
    nonempty = [line for line in lines if line.strip()]
    sample = nonempty[:50]
    delimiters = {
        "tab": sum("\t" in line for line in sample),
        "comma": sum("," in line for line in sample),
        "semicolon": sum(";" in line for line in sample),
        "whitespace": sum(bool(re.search(r"\S\s{2,}\S", line)) for line in sample),
    }
    return {
        "line_count": None if truncated else len(lines),
        "nonempty_line_count": None if truncated else len(nonempty),
        "sampled_line_count": len(lines),
        "sample_bytes": min(len(data), TEXT_PROFILE_BYTES),
        "scope": "prefix_only" if truncated else "complete_file",
        "likely_delimiter": max(delimiters, key=delimiters.get) if sample and max(delimiters.values()) else "unknown",
        "first_nonempty_line": clean(nonempty[0])[:300] if nonempty else "",
    }


def is_petrel_well_tops_ascii(path: Path) -> bool:
    if not looks_text(path):
        return False
    try:
        with path.open("rb") as handle:
            sample = handle.read(HEADER_SAMPLE_BYTES).decode("utf-8-sig", errors="replace")
    except OSError:
        return False
    upper = sample.upper()
    if "BEGIN HEADER" not in upper or "END HEADER" not in upper:
        return False
    header_match = re.search(r"BEGIN HEADER(.*?)END HEADER", sample, re.I | re.S)
    if not header_match:
        return False
    headers = {clean(line).lower() for line in header_match.group(1).splitlines() if clean(line)}
    return "well" in headers and "surface" in headers and bool(headers & {"md", "z", "twt picked", "twt auto"})


def convert_petrel_well_tops_ascii(source: Path, destination_root: Path) -> tuple[list[Path], dict[str, Any]]:
    lines = source.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    headers: list[str] = []
    rows: list[dict[str, Any]] = []
    in_header = False
    header_done = False
    version = ""
    comments: list[str] = []
    canonical_fields = [
        "record_class",
        "is_actual_pick_record",
        "petrel_export_confirmed",
        "native_binary_confirmed",
        "well_name",
        "surface",
        "x",
        "y",
        "depth",
        "measured_depth",
        "twt_picked",
        "twt_auto",
        "geological_age",
        "tvt",
        "tst",
        "interpreter",
        "observation_number",
        "dip_angle",
        "dip_azimuth",
        "missing",
        "confidence_factor",
        "used_by_depth_conversion",
        "used_by_geological_modeling",
        "symbol",
        "zone_log",
        "type",
        "source_file",
        "source_row",
        "decode_status",
    ]
    mapping = {
        "well_name": "Well",
        "surface": "Surface",
        "x": "X",
        "y": "Y",
        "depth": "Z",
        "measured_depth": "MD",
        "twt_picked": "TWT picked",
        "twt_auto": "TWT auto",
        "geological_age": "Geological age",
        "tvt": "TVT",
        "tst": "TST",
        "interpreter": "Interpreter",
        "observation_number": "Observation number",
        "dip_angle": "Dip angle",
        "dip_azimuth": "Dip azimuth",
        "missing": "Missing",
        "confidence_factor": "Confidence factor",
        "used_by_depth_conversion": "Used by dep.conv.",
        "used_by_geological_modeling": "Used by geo mod",
        "symbol": "Symbol",
        "zone_log": "Zone log",
        "type": "Type",
    }
    for line_number, raw in enumerate(lines, start=1):
        stripped = raw.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            comments.append(stripped)
            continue
        if stripped.upper().startswith("VERSION "):
            version = stripped
            continue
        if stripped.upper() == "BEGIN HEADER":
            in_header = True
            continue
        if stripped.upper() == "END HEADER":
            in_header = False
            header_done = True
            continue
        if in_header:
            headers.append(stripped)
            continue
        if not header_done:
            continue
        values = shlex.split(raw, posix=True)
        if len(values) != len(headers):
            raise ValueError(f"Line {line_number} has {len(values)} values; expected {len(headers)} from the Petrel header")
        raw_row = dict(zip(headers, values))
        row = {field: raw_row.get(source_field, "") for field, source_field in mapping.items()}
        row.update(
            {
                "record_class": "petrel_exported_well_tops_ascii",
                "is_actual_pick_record": "yes",
                "petrel_export_confirmed": "yes",
                "native_binary_confirmed": "no",
                "source_file": str(source),
                "source_row": line_number,
                "decode_status": "parsed_from_petrel_well_tops_ascii_companion",
            }
        )
        rows.append(row)

    destination_root.parent.mkdir(parents=True, exist_ok=True)
    csv_path = destination_root.with_suffix(".csv")
    summary_path = destination_root.with_suffix(".summary.json")
    generated: list[Path] = []
    if rows:
        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=canonical_fields)
            writer.writeheader()
            writer.writerows(rows)
        generated.append(csv_path)
    summary = {
        "converter": "portable_petrel_well_tops_ascii_parser",
        "source": str(source),
        "records": len(rows),
        "version": version,
        "source_header_columns": headers,
        "comments": comments,
        "petrel_export_confirmed": True,
        "native_binary_decode": False,
        "validation_boundary": "Rows were parsed from an open Petrel Well Tops ASCII companion. Confirm CRS, units, counts, and values against the Petrel-authored file.",
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    generated.append(summary_path)
    return generated, summary


def unique_headers(values: list[str]) -> list[str]:
    output: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = clean(value) or f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        output.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return output


def convert_las(source: Path, destination_root: Path) -> tuple[list[Path], dict[str, Any]]:
    import lasio  # type: ignore

    las = lasio.read(str(source))
    frame = las.df().reset_index()
    destination_root.parent.mkdir(parents=True, exist_ok=True)
    csv_path = destination_root.with_suffix(".csv")
    frame.to_csv(csv_path, index=False, float_format="%.8g")
    summary_path = destination_root.with_suffix(".summary.json")
    well = {item.mnemonic: clean(item.value) for item in las.well}
    summary = {
        "converter": "lasio",
        "source": str(source),
        "well_name": well.get("WELL", ""),
        "rows": int(len(frame)),
        "curves": [curve.mnemonic for curve in las.curves],
        "curve_units": {curve.mnemonic: curve.unit for curve in las.curves},
        "index_min": float(frame.iloc[:, 0].min()) if len(frame) else None,
        "index_max": float(frame.iloc[:, 0].max()) if len(frame) else None,
        "semantic_validation": "structure_and_row_depth_coverage_only",
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return [csv_path, summary_path], summary


def convert_workbook(source: Path, destination_root: Path, *,
                     max_seconds: float = WORKBOOK_MAX_SECONDS,
                     max_cells: int = WORKBOOK_MAX_CELLS) -> tuple[list[Path], dict[str, Any]]:
    from openpyxl import load_workbook  # type: ignore

    if not 0 < max_seconds < float('inf') or max_cells < 1:
        raise ValueError("Workbook conversion limits must be finite and positive")
    started = time.monotonic()
    workbook = None
    outputs: list[Path] = []
    sheets: list[dict[str, Any]] = []
    owned: list[Path] = []
    total_rows = total_cells = 0
    last_progress = started
    # Refuse reuse so cleanup touches only this invocation's new derivatives.
    destination_root.mkdir(parents=True, exist_ok=False)
    workbook_event('companion_workbook_started', source=str(source), output=str(destination_root),
                   max_seconds=max_seconds, max_cells=max_cells, buffer_bytes=WORKBOOK_BUFFER_BYTES)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
        used_names: set[str] = set()
        for index, sheet in enumerate(workbook.worksheets, 1):
            safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", sheet.title).strip("_") or "sheet"
            name = f"{safe}.csv"
            while name.casefold() in used_names:
                name = f"{index}_{name}"
            used_names.add(name.casefold())
            target = destination_root / name
            from geoviewer_output_paths import fits_file, windows_length, WORKBOOK_FILENAME_RESERVE
            if windows_length(name)>WORKBOOK_FILENAME_RESERVE or not fits_file(target):
                raise WorkbookLimitError('Workbook output path exceeds the qualified Windows path profile; original workbook is preserved.')
            row_count = nonempty_cell_count = max_columns = 0
            workbook_event('companion_workbook_sheet_started', source=str(source), sheet=sheet.title,
                           declared_rows=sheet.max_row, declared_columns=sheet.max_column)
            owned.append(target)
            with target.open("x", newline="", encoding="utf-8", buffering=WORKBOOK_BUFFER_BYTES) as handle:
                writer = csv.writer(handle)
                for row in sheet.iter_rows(values_only=True):
                    now = time.monotonic()
                    if now - started >= max_seconds or total_cells + len(row) > max_cells:
                        raise WorkbookLimitError(
                            f"Workbook CSV conversion budget reached after {total_rows} rows/{total_cells} cells "
                            f"({now-started:.1f}s); limits {max_seconds:g}s/{max_cells} cells. "
                            "Original workbook is preserved; no partial CSV is registered.")
                    values = ["" if value is None else value for value in row]
                    writer.writerow(values)
                    row_count += 1
                    total_rows += 1
                    total_cells += len(values)
                    nonempty_cell_count += sum(value != "" for value in values)
                    max_columns = max(max_columns, len(values))
                    if now - last_progress >= 5:
                        workbook_event('companion_workbook_progress', source=str(source), sheet=sheet.title,
                                       sheet_rows=row_count, rows=total_rows, cells=total_cells,
                                       elapsed_seconds=now-started)
                        last_progress = now
            if nonempty_cell_count:
                outputs.append(target)
                sheets.append({"sheet": sheet.title, "rows": row_count, "columns": max_columns, "output": target.name, "status": "converted"})
            else:
                target.unlink(missing_ok=True)
                sheets.append({"sheet": sheet.title, "rows": row_count, "columns": max_columns, "output": "", "status": "empty_not_exported"})
        summary_path = destination_root / "workbook.summary.json"
        owned.append(summary_path)
        summary = {"converter": "openpyxl", "source": str(source), "sheets": sheets,
                   "limits": {"max_seconds": max_seconds, "max_cells": max_cells,
                              "time_limit_scope": "cooperative checks between rows; not an OS I/O timeout"}}
        summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
        outputs.append(summary_path)
        workbook_event('companion_workbook_completed', source=str(source), rows=total_rows,
                       cells=total_cells, elapsed_seconds=time.monotonic()-started)
        return outputs, summary
    except BaseException as exc:
        cleanup_errors = []
        for path in owned:
            try:
                path.unlink(missing_ok=True)
            except OSError as cleanup:
                cleanup_errors.append(str(cleanup))
        try:
            destination_root.rmdir()
        except OSError as cleanup:
            cleanup_errors.append(str(cleanup))
        workbook_event('companion_workbook_limited' if isinstance(exc, WorkbookLimitError) else 'companion_workbook_failed',
                       source=str(source), error=str(exc), rows=total_rows, cells=total_cells,
                       elapsed_seconds=time.monotonic()-started, cleanup_errors=cleanup_errors)
        raise
    finally:
        if workbook is not None:
            workbook.close()


def _parts(points: list[Any], indexes: list[int]) -> list[list[Any]]:
    starts = list(indexes) + [len(points)]
    return [points[starts[i] : starts[i + 1]] for i in range(len(starts) - 1)]


def geometry_wkt(shape: Any) -> str:
    points = list(shape.points)
    shape_type = int(shape.shapeType)
    if shape_type in {1, 11, 21} and points:
        return f"POINT ({points[0][0]} {points[0][1]})"
    if shape_type in {8, 18, 28}:
        return "MULTIPOINT (" + ", ".join(f"({x} {y})" for x, y, *_ in points) + ")"
    parts = _parts(points, list(shape.parts)) if points else []
    if shape_type in {3, 13, 23}:
        body = ", ".join("(" + ", ".join(f"{x} {y}" for x, y, *_ in part) + ")" for part in parts)
        return ("LINESTRING " + body) if len(parts) == 1 else ("MULTILINESTRING (" + body + ")")
    if shape_type in {5, 15, 25, 31}:
        body = ", ".join("(" + ", ".join(f"{x} {y}" for x, y, *_ in part) + ")" for part in parts)
        return "POLYGON (" + body + ")"
    return ""


def convert_shapefile(source: Path, destination_root: Path) -> tuple[list[Path], dict[str, Any]]:
    import shapefile  # type: ignore

    # Missing attributes are a distinct scope, not a reason to lose readable
    # geometry. Open exact sidecars explicitly: permission/corruption failures
    # must not be swallowed by the reader and mistaken for missing DBF.
    source = Path(source)
    destination_root = Path(destination_root)
    targets = [destination_root.with_suffix(s) for s in ('.csv','.geojson','.prj','.summary.json')]
    if any(p.exists() or p.is_symlink() for p in targets):
        raise ValueError('Shapefile conversion destination already exists')
    moved = []
    with ExitStack() as stack:
        streams = {'shp': stack.enter_context(source.open('rb'))}
        for suffix in ('shx', 'dbf', 'cpg'):
            try:
                stream = source.with_suffix('.'+suffix).open('rb')
            except FileNotFoundError:
                stream = None
            streams[suffix] = stack.enter_context(stream) if stream is not None else None
        with shapefile.Reader(**streams) as reader:
            destination_root.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(prefix='shp_conversion_',dir=destination_root.parent) as pending:
                generated, summary = _convert_shapefile_reader(reader, source,
                    Path(pending)/destination_root.name, attributes_available=streams['dbf'] is not None)
                try:
                    for path in generated:
                        target = destination_root.parent/path.name
                        # Exclusive creation also works on USB filesystems without
                        # hardlinks. A failure removes only our new destinations.
                        with target.open('xb') as output:
                            moved.append(target)
                            with path.open('rb') as original:
                                shutil.copyfileobj(original,output,length=1024*1024)
                            output.flush()
                            os.fsync(output.fileno())
                except BaseException:
                    for path in moved:
                        path.unlink()
                    raise
    return moved, summary


def _convert_shapefile_reader(reader, source: Path, destination_root: Path, *, attributes_available: bool):
    import shapefile  # type: ignore

    fields = unique_headers([field[0] for field in reader.fields[1:]]) if attributes_available else []
    destination_root.parent.mkdir(parents=True, exist_ok=True)
    csv_path = destination_root.with_suffix(".csv")
    geojson_path = destination_root.with_suffix(".geojson")
    converted_prj_path = destination_root.with_suffix(".prj")
    prj_path = source.with_suffix(".prj")
    crs_wkt = prj_path.read_text(encoding="utf-8", errors="replace").strip() if prj_path.exists() else ""
    row_count = 0
    point_count = 0
    vertex_count = 0
    deleted_attributes = 0
    geometry_fields = [
        "source_shape_index",
        "source_attribute_status",
        "geometry_x",
        "geometry_y",
        "geometry_z",
        "geometry_m",
        "geometry_min_x",
        "geometry_min_y",
        "geometry_max_x",
        "geometry_max_y",
        "geometry_vertex_count",
        "geometry_wkt",
        "geometry_geojson",
        "shape_type",
        "source_crs_wkt",
    ]
    if set(fields).intersection(geometry_fields):
        raise ValueError('Shapefile DBF field conflicts with geometry/provenance column')
    absent = object()
    # Keep DBF deletion slots aligned with their original shape positions.
    # A short/extra attribute table is a failure, never a truncated zip result.
    records = (zip_longest(reader.iterShapes(),reader.iterRecords(deleted_as_None=True),fillvalue=absent)
               if attributes_available else ((shape,None) for shape in reader.iterShapes()))
    with csv_path.open("w", newline="", encoding="utf-8") as csv_handle, geojson_path.open("w", encoding="utf-8") as geojson_handle:
        writer = csv.DictWriter(csv_handle, fieldnames=fields + geometry_fields)
        writer.writeheader()
        geojson_handle.write('{"type":"FeatureCollection","features":[')
        first_feature = True
        for shape, record in records:
            if shape is absent or record is absent:
                raise ValueError('Shapefile geometry/attribute record count mismatch')
            points = list(shape.points)
            bbox = list(getattr(shape, "bbox", []) or [])
            geometry = None if int(shape.shapeType) == 0 else getattr(shape, "__geo_interface__", None)
            if attributes_available and record is not None and len(record) != len(fields):
                raise ValueError('Shapefile attribute field/value count mismatch')
            properties = dict(zip(fields, list(record))) if record is not None else {}
            attribute_status = ('dbf_record_read' if record is not None else
                                'dbf_record_marked_deleted' if attributes_available else 'missing_dbf')
            if attribute_status == 'dbf_record_marked_deleted':
                deleted_attributes += 1
            x = y = z = m = ""
            if int(shape.shapeType) in {1, 11, 21} and points:
                x, y = points[0][0], points[0][1]
                z_values = list(getattr(shape, "z", []) or [])
                m_values = list(getattr(shape, "m", []) or [])
                z = z_values[0] if z_values else ""
                m = m_values[0] if m_values else ""
                point_count += 1
            geometry_json = json.dumps(geometry, ensure_ascii=True, separators=(",", ":"), default=str, allow_nan=False) if geometry else ""
            row = dict(properties)
            row.update({
                "source_shape_index": row_count,
                "source_attribute_status": attribute_status,
                "geometry_x": x,
                "geometry_y": y,
                "geometry_z": z,
                "geometry_m": m,
                "geometry_min_x": bbox[0] if len(bbox) >= 4 else x,
                "geometry_min_y": bbox[1] if len(bbox) >= 4 else y,
                "geometry_max_x": bbox[2] if len(bbox) >= 4 else x,
                "geometry_max_y": bbox[3] if len(bbox) >= 4 else y,
                "geometry_vertex_count": len(points),
                "geometry_wkt": geometry_wkt(shape),
                "geometry_geojson": geometry_json,
                "shape_type": str(shape.shapeType),
                "source_crs_wkt": crs_wkt,
            })
            writer.writerow(row)
            feature = {"type": "Feature", "properties": properties, "geometry": geometry}
            if not first_feature:
                geojson_handle.write(",")
            geojson_handle.write(json.dumps(feature, ensure_ascii=True, separators=(",", ":"), default=str, allow_nan=False))
            first_feature = False
            row_count += 1
            vertex_count += len(points)
        geojson_handle.write("]}")

    generated: list[Path] = []
    if row_count:
        generated.extend([csv_path, geojson_path])
        if prj_path.exists():
            shutil.copy2(prj_path, converted_prj_path)
            generated.append(converted_prj_path)
    else:
        csv_path.unlink(missing_ok=True)
        geojson_path.unlink(missing_ok=True)
    summary_path = destination_root.with_suffix(".summary.json")
    summary = {
        "converter": "pyshp",
        "source": str(source),
        "records": row_count,
        "fields": fields,
        "attribute_status": ("dbf_records_read_with_deleted_rows" if deleted_attributes else "dbf_records_read") if attributes_available else "missing_dbf_geometry_only",
        "attributes_complete": attributes_available and deleted_attributes == 0,
        "deleted_dbf_records": deleted_attributes,
        "record_identity": "CSV source_shape_index and GeoJSON feature order are zero-based source SHP record order; deleted DBF slots retain geometry with empty properties",
        "shape_type": reader.shapeType,
        "shape_type_name": getattr(shapefile, "SHAPETYPE_LOOKUP", {}).get(reader.shapeType, "unknown"),
        "bbox": list(reader.bbox) if reader.bbox else [],
        "point_records_with_explicit_xy": point_count,
        "total_vertices": vertex_count,
        "csv_output": csv_path.name if row_count else "",
        "geojson_output": geojson_path.name if row_count else "",
        "prj_output": converted_prj_path.name if row_count and prj_path.exists() else "",
        "crs_status": "source_sidecar_preserved_not_semantically_resolved" if crs_wkt else "missing",
        "validation_boundary": "Geometry conversion is structural only. Missing DBF attributes are not recovered or inferred. Deleted DBF slots retain geometry but no attribute values. Confirm CRS, axis order, units, feature count, attributes, and topology against the source before use. GeoJSON/WKT are XY representations; explicit point Z/M are additionally retained in CSV. Other Z/M and multipatch semantics are not qualified.",
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    generated.append(summary_path)
    return generated, summary


def manifest_row(headers: list[str], export_package: Path, path: Path, project_name: str, petrel_version: str, source_type: str, notes: str) -> dict[str, str]:
    relative = rel(export_package, path)
    row = {header: "" for header in headers}
    row.update({
        "export_id": "portable_companion_" + hashlib.sha256(relative.lower().encode()).hexdigest()[:16],
        "project_name": project_name,
        "petrel_version": petrel_version,
        "export_date_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "source_object_path": "Portable companion extraction/" + relative,
        "source_object_type": source_type,
        "export_format": path.suffix.lstrip(".").upper() or "TEXT",
        "export_file": relative,
        "export_status": "exported_zero_gui_read_only",
        "validation_status": "unchecked",
        "sha256": sha256(path),
        "notes": notes,
    })
    return row


def upsert_manifest(path: Path, additions: list[dict[str, str]]) -> None:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        headers = list(reader.fieldnames or [])
        rows = list(reader)
    by_file = {row.get("export_file", "").lower(): row for row in rows}
    for row in additions:
        by_file[row["export_file"].lower()] = row
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(by_file.values())


def main() -> int:
    progress.enable_child_events()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-file", required=True)
    parser.add_argument("--export-package", required=True)
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--petrel-version", default="unknown")
    parser.add_argument("--mode", choices=["inventory", "copy", "convert"], default="convert")
    parser.add_argument("--max-file-bytes", type=int, default=2_000_000_000)
    parser.add_argument("--workbook-max-seconds", type=float, default=WORKBOOK_MAX_SECONDS)
    parser.add_argument("--workbook-max-cells", type=int, default=WORKBOOK_MAX_CELLS)
    parser.add_argument('--reference-seismic', action='store_true')
    args = parser.parse_args()
    if args.max_file_bytes < 1:
        parser.error("--max-file-bytes must be positive")
    if not 0 < args.workbook_max_seconds < float('inf') or args.workbook_max_cells < 1:
        parser.error("Workbook conversion limits must be finite and positive")

    project_file = Path(args.project_file).resolve()
    project_root = project_file.parent
    export_package = Path(args.export_package).resolve()
    ptd_root = project_root / f"{project_file.stem}.ptd"
    manifest_path = export_package / "00_manifest" / "export_manifest.csv"
    if not project_file.is_file() or not ptd_root.is_dir() or not manifest_path.is_file():
        raise SystemExit("Project file, matching PTD directory, or export manifest is missing")
    if is_relative_to(export_package, project_root):
        raise SystemExit("Export package must be outside the source project directory")

    # A source directory can contain more than one Petrel project. Never treat
    # a neighboring .pet file or any .ptd store as an ordinary companion: that
    # would mix native stores from different projects in one evidence package.
    from geoviewer_paths import project_files
    pruned_toolkit_roots=[]
    walked_files=list(project_files(project_root,pruned_toolkit_roots))
    petrel_project_files=sorted(path.resolve() for path in walked_files if path.suffix.lower()=='.pet')
    petrel_store_roots=sorted({parent.resolve() for path in walked_files for parent in path.parents
                              if parent.suffix.lower()=='.ptd' and is_relative_to(parent,project_root)})
    neighbor_project_files = [path for path in petrel_project_files if path != project_file]
    neighbor_store_roots = [path for path in petrel_store_roots if path != ptd_root]

    # The portable toolkit may intentionally be placed beside one or more
    # Petrel projects so its BAT launcher can discover them. Detect that exact
    # package signature and keep toolkit/runtime files out of project companion
    # inventories and conversions.
    launcher_file = project_root / "run_portable_petrel_extract.bat"
    toolkit_markers = [
        launcher_file,
        project_root / "toolkit.json",
        project_root / "scripts" / "invoke_portable_petrel_extract.ps1",
    ]
    co_located_toolkit_detected = all(path.is_file() for path in toolkit_markers)
    toolkit_directory_names = [".venv", ".agents", "scripts", "00_manifest", "runtime", "bootstrap", "build"]
    if launcher_file.is_file() and (project_root/'PetrelExtractor/STANDALONE.txt').is_file():
        co_located_toolkit_detected=True
        toolkit_directory_names=['PetrelExtractor']
    toolkit_file_names = [
        "run_portable_petrel_extract.bat",
        "README.md",
        "AGENTS.md",
        "requirements-core.txt",
        "requirements-geodata.txt",
        "toolkit.json",
        "STANDALONE.txt",
        "THIRD_PARTY_NOTICES.md",
        "requirements-standalone-lock.txt",
    ]
    toolkit_excluded_roots = (
        [(project_root / name).resolve() for name in toolkit_directory_names if (project_root / name).is_dir()]
        if co_located_toolkit_detected
        else []
    )
    toolkit_excluded_files = {launcher_file.resolve()} if launcher_file.is_file() else set()
    toolkit_excluded_roots.extend(path.resolve() for path in pruned_toolkit_roots)
    if pruned_toolkit_roots:co_located_toolkit_detected=True
    if co_located_toolkit_detected:
        toolkit_excluded_files.update(
            (project_root / name).resolve() for name in toolkit_file_names if (project_root / name).is_file()
        )

    optional = {}
    for module in ("lasio", "openpyxl", "shapefile"):
        try:
            __import__(module)
            optional[module] = "available"
        except ImportError:
            optional[module] = "unavailable"

    rows: list[dict[str, Any]] = []
    outputs: list[tuple[Path, str, str]] = []
    converted = 0
    preserved = 0
    failed = 0
    limited = 0
    skipped = 0
    toolkit_files_excluded = 0
    source_copy_root = export_package / "09_source_companions"
    converted_root = export_package / "10_converted_ascii"

    for source in sorted(walked_files):
        if not source.is_file():
            continue
        source_resolved = source.resolve()
        if source_resolved in petrel_project_files or any(
            is_relative_to(source_resolved, store_root) for store_root in petrel_store_roots
        ):
            continue
        if source_resolved in toolkit_excluded_files or any(
            is_relative_to(source_resolved, toolkit_root) for toolkit_root in toolkit_excluded_roots
        ):
            toolkit_files_excluded += 1
            continue
        if is_relative_to(source, export_package):
            continue
        source_rel = rel(project_root, source)
        size = source.stat().st_size
        over_size_limit = size > args.max_file_bytes
        category = classify(source)
        if args.reference_seismic and source.suffix.lower() in {'.zgy','.sgy','.segy'}:
            reference_row = dict.fromkeys(INVENTORY_FIELDS, "")
            reference_row.update({
                'source_relative_path': source_rel, 'category': category, 'size_bytes': size,
                'extension': source.suffix.lower(), 'status': 'referenced_seismic_not_copied_or_hashed',
                'text_profile_scope': 'not_text',
                'error': 'See project_seismic_inventory.json and the top-level report',
            })
            rows.append(reference_row)
            continue
        # This stage owns companion integrity; the native wrapper no longer
        # scans/hashes unrelated companions before native decoding can start.
        source_hash_before = sha256(source)
        if not over_size_limit and category in {"text_ascii", "text_ascii_no_extension", "ambiguous_dat_or_zmap"} and is_petrel_well_tops_ascii(source):
            category = "petrel_well_tops_ascii"
        status = "inventoried"
        preserved_path = ""
        converted_paths: list[str] = []
        conversion_gap = ""
        error = ""
        profile: dict[str, Any] = {"scope": "not_sampled_size_limit" if over_size_limit else "not_text"}
        text_like = not over_size_limit and looks_text(source)
        if text_like:
            try:
                profile = profile_text(source)
            except OSError as exc:
                error = str(exc)

        if args.mode in {"copy", "convert"}:
            if size == 0:
                status = "empty_source_not_copied"
                skipped += 1
            elif over_size_limit:
                status = "skipped_size_limit"
                skipped += 1
            elif source.suffix.lower() in PRESERVE_SUFFIXES or text_like:
                target = source_copy_root / Path(source_rel)
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, target)
                if source_hash_before != sha256(target):
                    raise RuntimeError(f"Copy hash mismatch: {source}")
                preserved_path = rel(export_package, target)
                outputs.append((target, "companion_source_preserved", "Hash-verified source companion copy."))
                preserved += 1
                status = "preserved"

        # A skipped/empty file must never reach a converter or acquire a false
        # "converted_and_preserved" status without a preserved source copy.
        if args.mode == "convert" and preserved_path:
            try:
                generated: list[Path] = []
                if category == "las_well_data" and optional["lasio"] == "available":
                    generated, _ = convert_las(source, converted_root / "las" / Path(source_rel))
                elif category == "excel_workbook" and optional["openpyxl"] == "available":
                    generated, _ = convert_workbook(source, converted_root / "workbooks" / Path(source_rel).with_suffix(""),
                                                    max_seconds=args.workbook_max_seconds,
                                                    max_cells=args.workbook_max_cells)
                elif category == "gis_shapefile_geometry" and optional["shapefile"] == "available":
                    generated, shape_summary = convert_shapefile(source, converted_root / "gis" / Path(source_rel))
                    if not shape_summary['attributes_complete']:
                        conversion_gap = shape_summary['attribute_status']
                elif category == "petrel_well_tops_ascii":
                    generated, _ = convert_petrel_well_tops_ascii(
                        source,
                        export_package / "02_wells" / "well_tops" / Path(source_rel),
                    )
                if generated:
                    converted += 1
                    status = "converted_and_preserved"
                    for target in generated:
                        converted_paths.append(rel(export_package, target))
                        outputs.append((target, "companion_converted_ascii", f"Converted from {source_rel}; source preserved separately."))
                elif category in {"las_well_data", "excel_workbook", "gis_shapefile_geometry", "petrel_well_tops_ascii"}:
                    status = "preserved_converter_unavailable"
            except WorkbookLimitError as exc:
                limited += 1
                status = "preserved_conversion_limited"
                error = str(exc)
            except Exception as exc:  # file-specific degradation is intentional
                failed += 1
                status = "preserved_conversion_failed"
                error = f"{type(exc).__name__}: {exc}"

        source_hash_after = sha256(source)
        if source_hash_after != source_hash_before:
            raise RuntimeError(f"Companion changed during processing: {source}")
        rows.append({
            "source_relative_path": source_rel,
            "category": category,
            "extension": source.suffix.lower(),
            "size_bytes": size,
            "sha256": source_hash_after,
            "text_line_count": profile.get("line_count", ""),
            "text_nonempty_line_count": profile.get("nonempty_line_count", ""),
            "text_profile_scope": profile.get("scope", ""),
            "text_sample_bytes": profile.get("sample_bytes", ""),
            "text_sample_line_count": profile.get("sampled_line_count", ""),
            "likely_delimiter": profile.get("likely_delimiter", ""),
            "first_nonempty_line": profile.get("first_nonempty_line", ""),
            "status": status,
            "preserved_file": preserved_path,
            "converted_files": ";".join(converted_paths),
            "conversion_gap": conversion_gap,
            "error": error,
        })

    inventory_path = export_package / "00_manifest" / "companion_source_inventory.csv"
    with inventory_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=INVENTORY_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    unsupported = [row for row in rows if row["category"] in {"unsupported_or_unknown", "legacy_excel_workbook", "ambiguous_dat_or_zmap", "segy_seismic", "zgy_seismic", "resqml_data", "resqml_epc_package"} or "failed" in row["status"] or "unavailable" in row["status"] or "limited" in row["status"] or row["status"] == "skipped_size_limit" or row["conversion_gap"]]
    unsupported_path = export_package / "99_unexported_or_manual" / "portable_unsupported_inventory.csv"
    if unsupported:
        unsupported_path.parent.mkdir(parents=True, exist_ok=True)
        with unsupported_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=INVENTORY_FIELDS)
            writer.writeheader()
            writer.writerows(unsupported)

    report_path = export_package / "07_workflows_reports" / "portable_extractor" / "companion_capability_report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    category_counts: dict[str, int] = {}
    for row in rows:
        category_counts[row["category"]] = category_counts.get(row["category"], 0) + 1
    report = {
        "operation": "portable_petrel_companion_extract",
        "created_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "source_project": str(project_file),
        "source_mutated": False,
        "petrel_launched": False,
        "mode": args.mode,
        "optional_dependencies": optional,
        "limits": {"max_preserved_file_bytes": args.max_file_bytes,
                   "workbook_max_seconds": args.workbook_max_seconds,
                   "workbook_max_cells": args.workbook_max_cells,
                   "workbook_time_scope": "cooperative checks between rows; blocked OS I/O needs the parent timeout",
                   "text_detection_bytes": TEXT_SAMPLE_BYTES,
                   "text_header_bytes": HEADER_SAMPLE_BYTES,
                   "text_profile_bytes": TEXT_PROFILE_BYTES},
        "counts": {
            "inventoried": len(rows),
            "preserved": preserved,
            "converted_source_files": converted,
            "converted_sources_with_gaps": sum(bool(row['conversion_gap']) for row in rows),
            "conversion_failures": failed,
            "conversion_limited": limited,
            "skipped_size_limit": skipped,
            "unsupported_or_manual": len(unsupported),
            "neighbor_petrel_project_files_excluded": len(neighbor_project_files),
            "neighbor_petrel_store_directories_excluded": len(neighbor_store_roots),
            "co_located_toolkit_files_excluded": toolkit_files_excluded,
        },
        "excluded_neighbor_petrel_projects": {
            "pet_files": [rel(project_root, path) for path in neighbor_project_files],
            "ptd_directories": [rel(project_root, path) for path in neighbor_store_roots],
            "reason": "Native stores belonging to other Petrel projects are not companion files and were not ingested.",
        },
        "co_located_toolkit_exclusion": {
            "detected": co_located_toolkit_detected,
            "launcher_detected": launcher_file.is_file(),
            "files_excluded": toolkit_files_excluded,
            "root_files": [rel(project_root, path) for path in sorted(toolkit_excluded_files)],
            "directories": [rel(project_root, path) for path in toolkit_excluded_roots],
            "reason": "Portable toolkit and runtime files are not Petrel project companion data.",
        },
        "category_counts": category_counts,
        "boundaries": [
            "Native proprietary arrays are not decoded by this companion-file stage.",
            "Preservation and structural conversion do not establish CRS, scientific validity, or semantic completeness.",
            "SEG-Y, ZGY, ZMAP, RESQML, and ambiguous DAT files require their specialist validators.",
            "Files above the size limit are inventoried and hashed in chunks, without copying or conversion.",
            "Text profiles marked prefix_only contain sample counts, not full-file line counts.",
            "Workbook conversion limits preserve the original workbook and omit incomplete CSV derivatives; they never limit native binary recovery.",
        ],
    }
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")

    outputs.extend([
        (inventory_path, "companion_source_inventory", "Source companion inventory with hashes and conversion status."),
        (report_path, "portable_extraction_report", "Capability and degradation report."),
    ])
    if unsupported:
        outputs.append((unsupported_path, "unsupported_or_manual_inventory", "Files requiring specialist validation or manual handling."))
    with manifest_path.open(newline="", encoding="utf-8-sig") as handle:
        headers = list(csv.DictReader(handle).fieldnames or [])
    additions = [manifest_row(headers, export_package, path, args.project_name, args.petrel_version, source_type, notes) for path, source_type, notes in outputs]
    upsert_manifest(manifest_path, additions)

    print("Portable companion extraction: completed")
    print(f"Inventory: {inventory_path}")
    print(f"Capability report: {report_path}")
    print(f"Inventoried: {len(rows)}")
    print(f"Preserved: {preserved}")
    print(f"Converted source files: {converted}")
    print(f"Conversion failures: {failed}")
    print(f"Workbook conversions limited (original preserved): {limited}")
    print(f"Neighbor Petrel project files excluded: {len(neighbor_project_files)}")
    print(f"Neighbor Petrel store directories excluded: {len(neighbor_store_roots)}")
    print(f"Co-located portable toolkit files excluded: {toolkit_files_excluded}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
